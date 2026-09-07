"""Pure gates + send-list materialization for ops.py execute.

No network. No broker sends.
Send SoT after materialize: execute_send_list.json
"""
from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from ops.ops_us_policy import is_integer_qty, normalize_symbol
from ops.ops_live_safety import validate_order_gates, validate_unique_rows

EPS_CASH = 1.0
MAX_NAMES = 40
MAX_NAME_W = 0.15
QTY_EPS = 1e-9
NOTIONAL_TOL = 1.0  # USD (or quote currency notional)
QTY_TOL = 1e-6


def zcode(x) -> str:
    """Normalize an order code. Numeric (KR legacy) codes keep 6-digit padding;
    anything else is treated as a US ticker via normalize_symbol."""
    s = str(x).replace(".0", "").strip()
    if s.isdigit():
        return s.zfill(6)
    return normalize_symbol(s)


@dataclass
class GateResult:
    ok: bool
    reasons: List[str]
    warnings: List[str]
    sendable: Any
    n_sell: int
    n_buy: int
    buy_notional: float
    sell_notional: float
    cash_usd: Optional[float]
    capital_usd: Optional[float]
    shortfall_net: Optional[float]
    health_status: str


def load_orders_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        raise FileNotFoundError(path)
    df = pd.read_csv(path)
    if df.empty:
        return df
    if "code" in df.columns:
        df["code"] = df["code"].map(zcode)
    if "side" in df.columns:
        df["side"] = df["side"].astype(str).str.upper()
    return df


def load_trade_sheet_rows(xlsx_path: Path) -> pd.DataFrame:
    """Parse 05_Trade order ticket table (header row with prio/code/side/qty)."""
    if not xlsx_path.exists():
        raise FileNotFoundError(xlsx_path)
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path, data_only=True, read_only=True)
    if "05_Trade" not in wb.sheetnames:
        wb.close()
        raise ValueError("05_Trade sheet missing")
    ws = wb["05_Trade"]
    header_row = None
    headers: List[str] = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=40, values_only=True), 1):
        vals = ["" if c is None else str(c).strip() for c in (row[:20] if row else ())]
        if "code" in vals and "side" in vals and ("qty" in vals or "notional" in vals):
            header_row = i
            headers = vals
            break
    if header_row is None:
        wb.close()
        raise ValueError("05_Trade order header not found")
    idx = {h: n for n, h in enumerate(headers) if h}
    rows: List[Dict[str, Any]] = []
    for row in ws.iter_rows(min_row=header_row + 1, max_row=header_row + 200, values_only=True):
        if (not row) or (row[idx.get("code", 1)] is None):
            if all((c is None or str(c).strip() == "") for c in (row or ())[:5]):
                break
            continue
        code = zcode(row[idx["code"]])
        if not code:
            continue

        def _get(name: str, default=None):
            j = idx.get(name)
            if j is None:
                return default
            return row[j] if j < len(row) else default

        side = str(_get("side") or "").upper()
        try:
            qty = float(_get("qty") or 0)
        except (TypeError, ValueError):
            qty = 0.0
        try:
            notional = float(_get("notional") or 0)
        except (TypeError, ValueError):
            notional = 0.0
        try:
            if _get("est_px") is None:
                est_px = None
            else:
                est_px = float(_get("est_px") or 0)
        except (TypeError, ValueError):
            est_px = None
        rows.append(
            {
                "code": code,
                "name": _get("name") or "",
                "side": side,
                "action": _get("action") or "",
                "qty": qty,
                "notional": notional,
                "est_px": est_px,
                "prio": _get("prio"),
            }
        )
    wb.close()
    return pd.DataFrame(rows)


def cross_check(sheet: pd.DataFrame, csv: pd.DataFrame) -> List[str]:
    """Return list of mismatch reasons (empty = ok)."""
    reasons: List[str] = []
    if sheet is None or sheet.empty:
        reasons.append("sheet_empty")
        return reasons
    if csv is None or csv.empty:
        reasons.append("csv_empty")
        return reasons
    s = sheet.copy()
    c = csv.copy()
    s["code"] = s["code"].map(zcode)
    c["code"] = c["code"].map(zcode)
    s["side"] = s["side"].astype(str).str.upper()
    c["side"] = c["side"].astype(str).str.upper()
    s_pairs = [(r.code, r.side) for r in s.itertuples() if r.side in ("BUY", "SELL")]
    c_pairs = [(r.code, str(getattr(r, "side", "") or "").upper()) for r in c.itertuples()
               if str(getattr(r, "side", "") or "").upper() in ("BUY", "SELL")]
    if len(s_pairs) != len(set(s_pairs)):
        reasons.append("sheet_duplicate_order_keys")
    if len(c_pairs) != len(set(c_pairs)):
        reasons.append("csv_duplicate_order_keys")
    s_keys = set(s_pairs)
    c_keys = {
        (r.code, str(getattr(r, "side", "") or "").upper())
        for r in c.itertuples()
        if str(getattr(r, "side", "") or "").upper() in ("BUY", "SELL")
    }
    if s_keys != c_keys:
        only_s = s_keys - c_keys
        only_c = c_keys - s_keys
        if only_s:
            reasons.append(f"sheet_only_keys={sorted(list(only_s))[:5]}")
        if only_c:
            reasons.append(f"csv_only_keys={sorted(list(only_c))[:5]}")
    c_map: Dict[Tuple[str, str], Any] = {}
    for r in c.itertuples():
        side = str(getattr(r, "side", "") or "").upper()
        if side not in ("BUY", "SELL"):
            continue
        c_map[(zcode(r.code), side)] = r
    for r in s.itertuples():
        if r.side not in ("BUY", "SELL"):
            continue
        key = (zcode(r.code), r.side)
        o = c_map.get(key)
        if o is None:
            continue
        try:
            sq = float(r.qty or 0)
            cq = float(getattr(o, "qty", 0) or 0)
            if abs(sq - cq) > QTY_TOL:
                reasons.append(f"qty_mismatch {key} sheet={sq} csv={cq}")
        except (TypeError, ValueError):
            reasons.append(f"qty_parse {key}")
        try:
            sn = float(r.notional or 0)
            cn = float(getattr(o, "notional", 0) or 0)
            if abs(sn - cn) > NOTIONAL_TOL:
                reasons.append(f"notional_mismatch {key} sheet={sn} csv={cn}")
        except (TypeError, ValueError):
            reasons.append(f"notional_parse {key}")
    return reasons


def holdings_qty_map(book: Optional[Dict[str, Any]]) -> Dict[str, float]:
    out: Dict[str, float] = {}
    if not book:
        return out
    rows = book.get("us_holdings") or []
    for h in rows:
        raw = h.get("symbol_raw") or h.get("code") or h.get("symbol") or ""
        code = zcode(raw)
        if not code:
            continue
        try:
            q = float(h.get("qty") or h.get("quantity") or 0)
        except (TypeError, ValueError):
            q = 0.0
        out[code] = out.get(code, 0.0) + q
    return out


def materialize_send_list(
    orders: pd.DataFrame,
    *,
    holdings: Optional[Dict[str, float]] = None,
    capital_usd: Optional[float] = None,
) -> List[Dict[str, Any]]:
    """Build send rows: HOLD skipped; SELL qty from live holdings when available; recompute notional."""
    holdings = holdings or {}
    rows: List[Dict[str, Any]] = []
    if orders is None or orders.empty:
        return rows
    df = orders.copy()
    if "side" in df.columns:
        df["side"] = df["side"].astype(str).str.upper()
    if "code" in df.columns:
        df["code"] = df["code"].map(zcode)
    side_prio = {"SELL": 0, "BUY": 1}
    df["_sp"] = df["side"].map(lambda x: side_prio.get(str(x), 9))
    if "prio" in df.columns:
        df = df.sort_values(["_sp", "prio"], kind="mergesort")
    else:
        df = df.sort_values(["_sp", "code"], kind="mergesort")
    for r in df.itertuples():
        side = str(getattr(r, "side", "") or "").upper()
        if side not in ("BUY", "SELL"):
            continue
        code = zcode(getattr(r, "code", "") or "")
        raw_q = getattr(r, "qty", None)
        if raw_q is None or (isinstance(raw_q, float) and raw_q != raw_q):
            raw_q = getattr(r, "target_qty", None)
        try:
            if raw_q is None or (isinstance(raw_q, float) and raw_q != raw_q):
                qty = 0.0
            else:
                qty = float(raw_q)
        except (TypeError, ValueError):
            qty = 0.0
        est_px = getattr(r, "est_px", None)
        try:
            est_px = float(est_px) if est_px not in (None, "") else None
        except (TypeError, ValueError):
            est_px = None
        if side == "SELL":
            live_q = holdings.get(code)
            if live_q is not None and live_q > 0:
                if qty > 0:
                    if float(qty) <= float(live_q) + 1e-9:
                        qty = min(float(live_q), float(qty))
                    else:
                        qty = float(live_q)
                else:
                    qty = float(live_q)
        if qty <= 0:
            continue
        notional = None
        if est_px and est_px > 0:
            notional = float(qty) * float(est_px)
        else:
            try:
                notional = float(getattr(r, "notional", 0) or 0)
            except (TypeError, ValueError):
                notional = 0.0
        rows.append(
            {
                "code": code,
                "name": getattr(r, "name", "") or "",
                "side": side,
                "action": getattr(r, "action", "") or "",
                "send_qty": float(qty),
                "est_px": est_px,
                "notional": float(notional or 0.0),
                "capital_usd": capital_usd,
            }
        )
    return rows


def evaluate_gates(
    sendable: List[Dict[str, Any]],
    health: Dict[str, Any],
    *,
    phase: str = "A",
    cash_usd: Optional[float] = None,
    post_sell_cash_usd: Optional[float] = None,
    confirm_high_value: bool = False,
) -> GateResult:
    """Phase A: pre-send. Phase B: after all SELL full-fill, before BUY."""
    reasons: List[str] = []
    warnings: List[str] = []
    status = str(health.get("status") or "")
    sizing = health.get("sizing") or {}
    capital = health.get("capital_usd")
    if capital is None:
        capital = sizing.get("capital_usd")
    try:
        capital_f = float(capital) if capital is not None else None
    except (TypeError, ValueError):
        capital_f = None

    cash0 = cash_usd
    if cash0 is None:
        cash0 = sizing.get("cash_usd")
    if cash0 is None and isinstance(health.get("toss"), dict):
        cash0 = health["toss"].get("cash_usd")
    try:
        cash0_f = float(cash0) if cash0 is not None else None
    except (TypeError, ValueError):
        cash0_f = None

    if status == "BLOCK":
        reasons.append("health_BLOCK")

    reasons.extend(validate_unique_rows(sendable))

    sells = [r for r in sendable if r.get("side") == "SELL"]
    buys = [r for r in sendable if r.get("side") == "BUY"]

    seen_buy = False
    for r in sendable:
        if r.get("side") == "BUY":
            seen_buy = True
            continue
        if r.get("side") == "SELL" and seen_buy:
            reasons.append("sell_not_before_buy")
            break

    names = {r.get("code") for r in sendable}
    if len(names) > MAX_NAMES:
        reasons.append(f"n_names>{MAX_NAMES}:{len(names)}")

    sell_notional = sum(float(r.get("notional") or 0) for r in sells)
    buy_notional = sum(float(r.get("notional") or 0) for r in buys)

    for r in sendable:
        code = str(r.get("code") or "")
        qty = r.get("send_qty")
        if code.isdigit():
            if not is_integer_qty(qty):
                reasons.append(f"non_integer_qty:{code}:{qty}")
            continue
        try:
            if float(qty) <= 0:
                reasons.append(f"non_positive_qty:{code}:{qty}")
        except (TypeError, ValueError):
            reasons.append(f"bad_qty:{code}:{qty}")

    if capital_f and capital_f > 0:
        for r in buys:
            n = float(r.get("notional") or 0)
            if n > capital_f * MAX_NAME_W + EPS_CASH:
                reasons.append(f"name_cap15:{r.get('code')}:{n}")

    shortfall_net = None
    if phase == "A":
        if cash0_f is not None:
            if sells:
                shortfall_net = max(0.0, buy_notional - sell_notional - cash0_f)
                if shortfall_net > EPS_CASH:
                    reasons.append(f"shortfall_net>{EPS_CASH}:{shortfall_net}")
            else:
                if buy_notional > cash0_f + EPS_CASH:
                    reasons.append(f"buy_notional>cash:{buy_notional}>{cash0_f}")
                shortfall_net = max(0.0, buy_notional - cash0_f)
        else:
            warnings.append("cash_usd_missing")
    elif phase == "B":
        cash1 = post_sell_cash_usd
        try:
            cash1_f = float(cash1) if cash1 is not None else None
        except (TypeError, ValueError):
            cash1_f = None
        if cash1_f is None:
            reasons.append("post_sell_cash_missing")
        else:
            if buy_notional > cash1_f + EPS_CASH:
                reasons.append(f"BLOCKED_BUY_CASH:{buy_notional}>{cash1_f}")
            shortfall_net = max(0.0, buy_notional - cash1_f)
        sendable = buys

    safety_cash = post_sell_cash_usd if phase == "B" else cash0_f
    extra_reasons, extra_warnings = validate_order_gates(
        sendable if phase == "A" else buys,
        capital_usd=capital_f,
        cash_usd=safety_cash,
        phase=phase,
        confirm_high_value=confirm_high_value,
    )
    reasons.extend(extra_reasons)
    warnings.extend(extra_warnings)

    ok = len(reasons) == 0
    return GateResult(
        ok=ok,
        reasons=reasons,
        warnings=warnings,
        sendable=sendable,
        n_sell=len(sells),
        n_buy=len(buys),
        buy_notional=buy_notional,
        sell_notional=sell_notional,
        cash_usd=cash0_f,
        capital_usd=capital_f,
        shortfall_net=shortfall_net,
        health_status=status,
    )


def write_send_list(run_dir: Path, rows: List[Dict[str, Any]], meta: Optional[Dict[str, Any]] = None) -> Path:
    run_dir = Path(run_dir)
    payload = {"meta": meta or {}, "rows": rows}
    path = run_dir / "execute_send_list.json"
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def load_send_list(run_dir: Path) -> List[Dict[str, Any]]:
    path = Path(run_dir) / "execute_send_list.json"
    if not path.exists():
        return []
    data = json.loads(path.read_text(encoding="utf-8"))
    return list(data.get("rows") or [])


def gate_result_dict(g: GateResult) -> Dict[str, Any]:
    return asdict(g)
