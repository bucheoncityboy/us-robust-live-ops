"""ops.py execute orchestration - rebuild ticket, gates, LIVE Toss send.

Production: LIVE_ORDERS_ENABLED=True after research lock / human approval
(real order HTTP POSTs). If False, gates run then BLOCK with zero order HTTP.
"""

from __future__ import annotations

import json
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import pandas as pd

from ops.ops_live_safety import (
    acquire_execution_claim,
    execution_evidence,
    validate_account_identity,
    validate_execute_package,
    validate_official_next_open,
    write_attempt_result,
)
from ops.ops_execute_gates import (
    cross_check,
    evaluate_gates,
    gate_result_dict,
    holdings_qty_map,
    load_orders_csv,
    load_trade_sheet_rows,
    materialize_send_list,
    write_send_list,
)
from ops.toss_orders import (
    LIVE_ORDERS_ENABLED,
    LiveOrdersBlocked,
    place_and_await,
    research_block_reason,
    utc_now_iso,
)


def _find_xlsx(run_dir: Path) -> Optional[Path]:
    """Return the newest workbook by modification time, including recovery files."""
    xs = list(Path(run_dir).glob("US_Robust_Ops_*.xlsx"))
    return max(xs, key=lambda p: (p.stat().st_mtime_ns, p.name)) if xs else None


def _append_fill(path: Path, row: Dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("a", encoding="utf-8") as f:
        f.write(json.dumps(row, ensure_ascii=False) + "\n")


def apply_execute_status_merge(run_dir: Path, status_map: Dict[str, str]) -> Path:
    """Persist durable status map; excel rebuild must re-merge via this file."""
    run_dir = Path(run_dir)
    path = run_dir / "execute_status.json"
    prev = {}
    if path.exists():
        try:
            prev = json.loads(path.read_text(encoding="utf-8"))
        except Exception:
            prev = {}
    codes = prev.get("codes") if isinstance(prev.get("codes"), dict) else {}
    codes.update(status_map)
    payload = {
        "updated_utc": utc_now_iso(),
        "codes": codes,
    }
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return path


def patch_trade_status_column(xlsx_path: Path, status_map: Dict[str, str]) -> None:
    """Best-effort patch of 05_Trade status cells by code."""
    if not xlsx_path.exists() or not status_map:
        return
    from openpyxl import load_workbook

    wb = load_workbook(xlsx_path)
    if "05_Trade" not in wb.sheetnames:
        wb.close()
        return
    ws = wb["05_Trade"]
    header_row = None
    headers = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=40, values_only=True), 1):
        vals = [str(c).strip() if c is not None else "" for c in row[:20]]
        if "code" in vals and "side" in vals:
            header_row = i
            headers = vals
            break
    if header_row is None:
        wb.close()
        return
    idx = {h: n + 1 for n, h in enumerate(headers) if h}  # 1-based
    code_c = idx.get("code")
    status_c = idx.get("status")
    if not code_c or not status_c:
        wb.close()
        return
    for r in range(header_row + 1, header_row + 200):
        code = ws.cell(r, code_c).value
        if code is None:
            break
        from ops.ops_us_policy import normalize_symbol as _ns
        raw = str(code).replace(".0", "")
        code_s = raw.zfill(6) if raw.isdigit() else _ns(raw)
        if code_s in status_map:
            ws.cell(r, status_c).value = status_map[code_s]
    wb.save(xlsx_path)
    wb.close()



def _refresh_excel_after_execution(
    run_dir: Path,
    *,
    positions: str,
    prior_book: Optional[Dict[str, Any]],
    summary: Dict[str, Any],
    no_xlsx: bool,
) -> Optional[Path]:
    """Re-pull post-trade holdings and rebuild all live Excel sections on every order outcome."""
    if no_xlsx:
        return None
    book = prior_book
    if str(positions).strip().lower() in {"toss", "toss://", "toss:live"}:
        try:
            from ops.toss_portfolio import fetch_portfolio
            book = fetch_portfolio()
            book["asof_utc"] = utc_now_iso()
            summary["post_trade_portfolio"] = "FRESH"
        except Exception as exc:
            book = dict(prior_book or {})
            book["_stale"] = True
            book["status"] = "stale"
            book["ops_ready"] = False
            book["error"] = f"post-trade refresh failed: {type(exc).__name__}: {exc}"
            book.setdefault("us_holdings", list((prior_book or {}).get("us_holdings") or []))
            summary["post_trade_portfolio"] = "STALE / REFRESH ERROR"
            summary["post_trade_portfolio_error"] = book["error"]
    try:
        from ops.ops_excel_write import write_ops_xlsx
        out = write_ops_xlsx(Path(run_dir), positions=positions, toss_book=book)
        summary["excel"] = str(out)
        return out
    except Exception as exc:
        summary["excel_refresh_error"] = f"{type(exc).__name__}:{exc}"
        return None



def run_execute(
    run_dir: Path,
    *,
    positions: str = "toss",
    capital_usd: Optional[float] = None,
    asof: Optional[str] = None,
    no_xlsx: bool = False,
    risk_mode: Optional[str] = None,
    confirm_high_value: bool = False,
    confirm_outside_window: bool = False,
) -> Tuple[int, Dict[str, Any]]:
    """Run a package once. All mutating broker calls occur after fail-closed preflight."""
    run_dir = Path(run_dir)
    summary: Dict[str, Any] = {
        "run_dir": str(run_dir), "live_orders_enabled": bool(LIVE_ORDERS_ENABLED),
        "http_order_posts": 0, "phase": "package_validation", "risk_mode": risk_mode,
        "confirm_outside_window": bool(confirm_outside_window),
    }

    package = validate_execute_package(run_dir, asof=asof)
    summary["package_validation"] = package
    if not package.get("ok"):
        summary["error"] = "package_validation_fail"
        try:
            existing = sorted(
                p.name for p in run_dir.parent.iterdir()
                if p.is_dir() and (p / "target.csv").exists()
            )
        except OSError:
            existing = []
        if existing:
            print(f"  existing month packages: {', '.join(existing)}")
            print(f"  hint: use --run-dir results/ops_runs/{existing[-1]} to target the latest package")
        return 2, summary
    package_id = str(package.get("package_id") or run_dir.name)

    if LIVE_ORDERS_ENABLED:
        if not risk_mode:
            summary["error"] = "risk_mode_required"
            return 2, summary
        if str(positions).strip().lower() not in {"toss", "toss://", "toss:live"}:
            summary["error"] = "LIVE_requires_positions_toss"
            return 2, summary
        if capital_usd is not None:
            summary["error"] = "LIVE_capital_override_forbidden"
            return 2, summary
        from ops.ops_env import get_secret, load_env
        load_env()
        if not get_secret("TOSS_ACCOUNT_SEQ", "").strip():
            summary["error"] = "TOSS_ACCOUNT_SEQ_pin_required_for_LIVE"
            return 2, summary
    effective_risk = risk_mode or "NORMAL"

    summary["phase"] = "rebuild"
    from ops.ops_monthly_run import rebuild_rebalance_ticket, load_positions
    try:
        ticket = rebuild_rebalance_ticket(
            run_dir, positions=positions, capital_usd=capital_usd, risk_mode=effective_risk,
        )
    except Exception as exc:
        summary["error"] = f"rebuild_fail:{type(exc).__name__}:{exc}"
        return 2, summary
    book = getattr(load_positions, "last_toss_book", None)
    summary["ticket"] = {k: ticket.get(k) for k in ("n_buy", "n_sell", "capital_usd", "cash_usd", "risk_mode")}

    xpath = None
    if not no_xlsx:
        summary["phase"] = "excel"
        try:
            from ops.ops_excel_write import write_ops_xlsx
            xpath = write_ops_xlsx(run_dir, positions=positions, toss_book=book)
            summary["excel"] = str(xpath)
        except Exception as exc:
            summary["error"] = f"excel_build_fail:{type(exc).__name__}:{exc}"
            return 2, summary

    # CSV/package is always authoritative. Excel is comparison-only.
    summary["phase"] = "xcheck"
    csv_df = load_orders_csv(run_dir / "orders_preview.csv")
    if xpath and xpath.exists():
        try:
            sheet_df = load_trade_sheet_rows(xpath)
            mismatches = cross_check(sheet_df, csv_df)
        except Exception as exc:
            mismatches = [f"excel_compare_failed:{type(exc).__name__}:{exc}"]
        if mismatches:
            summary["xcheck_fail"] = mismatches
            return 2, summary

    health = json.loads((run_dir / "health.json").read_text(encoding="utf-8"))
    holdings = holdings_qty_map(book)
    capital = health.get("capital_usd")
    sendable = materialize_send_list(csv_df, holdings=holdings, capital_usd=capital)
    for row in sendable:
        row.setdefault("market", "US")
    write_send_list(run_dir, sendable, meta={"source": "orders_preview.csv", "package_id": package_id})
    summary["n_sendable"] = len(sendable)

    replay = execution_evidence(run_dir, sendable)
    summary["replay"] = replay
    if replay.get("blocked"):
        summary["phase"] = "replay_check"
        summary["error"] = "prior_or_ambiguous_execution_evidence"
        return 3, summary

    cash0 = None
    if book is not None:
        try:
            cash0 = float(book.get("cash_buying_power_usd") if book.get("cash_buying_power_usd") is not None else book.get("cash_usd"))
        except (TypeError, ValueError):
            cash0 = None
    summary["phase"] = "gates_A"
    gate_a = evaluate_gates(
        sendable, health, phase="A", cash_usd=cash0,
        confirm_high_value=confirm_high_value,
    )
    summary["gate_a"] = gate_result_dict(gate_a)
    if not gate_a.ok:
        summary["error"] = "gate_A_fail"
        return 3, summary
    if not sendable:
        summary["phase"] = "empty_success"
        summary["message"] = "no SELL/BUY rows to send"
        return 0, summary

    summary["phase"] = "live_check"
    block = research_block_reason()
    if block or not LIVE_ORDERS_ENABLED:
        summary["error"] = "LIVE_BLOCKED"
        summary["block_reason"] = block or "LIVE_ORDERS_ENABLED=False"
        from ops.toss_orders import dry_run_orders
        dry = dry_run_orders(sendable, market="US", run_tag=package_id)
        dry_path = run_dir / "execute_dry_run.json"
        dry_path.write_text(json.dumps(dry, ensure_ascii=False, indent=2), encoding="utf-8")
        summary["dry_run"] = str(dry_path)
        summary["dry_run_n_ok"] = dry.get("n_ok")
        summary["dry_run_n_err"] = dry.get("n_err")
        return 4, summary

    try:
        claim = acquire_execution_claim(run_dir, {
            "package_id": package_id, "risk_mode": effective_risk,
            "send_codes": [r.get("code") for r in sendable], "phase": "CONNECTING",
        })
        summary["claim"] = str(claim)
    except FileExistsError as exc:
        summary["phase"] = "claim"
        summary["error"] = str(exc)
        return 3, summary

    from ops.toss_orders import OrderClient
    summary["phase"] = "connect"
    try:
        client = OrderClient.connect()
    except Exception as exc:
        summary["error"] = f"connect_fail:{type(exc).__name__}:{exc}"
        summary["http_order_posts"] = 0
        write_attempt_result(run_dir, {
            "status": "CONNECT_FAIL_ZERO_POST", "http_order_posts": 0,
            "error": summary["error"], "package_id": package_id,
        })
        return 5, summary

    # Read-only broker preflight after connection, before any create-order POST.
    summary["phase"] = "broker_preflight"
    preflight_reasons = validate_account_identity(book, client)
    if not str(__import__('os').environ.get("TOSS_ACCOUNT_SEQ") or "").strip():
        preflight_reasons.append("TOSS_ACCOUNT_SEQ_pin_required_for_LIVE")
    try:
        calendar = client.get_us_market_calendar(str(health.get("signal_date")))
        window_reasons = validate_official_next_open(calendar, str(health.get("signal_date")))
        if window_reasons:
            if confirm_outside_window:
                summary["outside_window_confirmed"] = window_reasons
                window_reasons = []
            else:
                print("  !!! outside regular next-open window (use --confirm-outside-window to override explicitly)")
        preflight_reasons.extend(window_reasons)
        open_orders = client.get_open_orders()
        if open_orders:
            preflight_reasons.append(f"existing_open_orders:{len(open_orders)}")
        for row in sendable:
            if row.get("side") != "SELL":
                continue
            sellable = float(client.get_sellable_quantity(row["code"]))
            if float(row.get("send_qty") or 0) > sellable + 1e-9:
                preflight_reasons.append(
                    f"sellable_shortfall:{row['code']}:{row.get('send_qty')}>{sellable}"
                )
    except Exception as exc:
        preflight_reasons.append(f"broker_preflight_failed:{type(exc).__name__}:{exc}")
    if preflight_reasons:
        summary["preflight_reasons"] = preflight_reasons
        summary["error"] = "broker_preflight_fail"
        summary["http_order_posts"] = getattr(client, "http_post_count", 0)
        write_attempt_result(run_dir, {
            "status": "PREFLIGHT_FAIL_ZERO_POST", "http_order_posts": summary["http_order_posts"],
            "reasons": preflight_reasons, "package_id": package_id,
        })
        return 5, summary

    ts = datetime.now().strftime("%Y%m%d_%H%M%S")
    fills_path = run_dir / f"fills_{ts}.jsonl"
    status_map: Dict[str, str] = {}
    sells = [r for r in sendable if r["side"] == "SELL"]
    buys = [r for r in sendable if r["side"] == "BUY"]

    def persist_failure(code: int) -> Tuple[int, Dict[str, Any]]:
        summary["http_order_posts"] = getattr(client, "http_post_count", 0)
        summary["fills"] = str(fills_path)
        apply_execute_status_merge(run_dir, status_map)
        posts = int(summary["http_order_posts"] or 0)
        # Zero-POST failures (pre-send gates) are safe to auto-retry per the
        # documented rule; anything that touched the order endpoint needs review.
        status = "GATE_FAIL_ZERO_POST" if posts == 0 else "ORDER_ATTEMPT_REVIEW_REQUIRED"
        write_attempt_result(run_dir, {
            "status": status, "http_order_posts": posts,
            "error": summary.get("error"), "package_id": package_id,
        })
        _refresh_excel_after_execution(
            run_dir, positions=positions, prior_book=book, summary=summary, no_xlsx=no_xlsx,
        )
        return code, summary

    def send_batch(batch: List[Dict[str, Any]], label: str) -> bool:
        for index, row in enumerate(batch, 1):
            code, side = row["code"], row["side"]
            qty = float(row["send_qty"])
            try:
                att = place_and_await(
                    client, code=code, side=side, qty=qty,
                    run_tag=package_id, attempt=index,
                    order_amount=float(row.get("notional") or 0),
                    confirm_high_value=confirm_high_value,
                )
            except Exception as exc:
                _append_fill(fills_path, {"event": "ERROR", "utc": utc_now_iso(), "code": code,
                    "side": side, "qty": qty, "error": str(exc), "phase": label})
                status_map[code] = f"FAIL:{exc}"
                summary["error"] = f"{label}_fail:{code}:{exc}"
                return False
            _append_fill(fills_path, {
                "event": "FILL", "utc": utc_now_iso(), "code": code, "side": side, "qty": qty,
                "client_order_id": att.client_order_id, "order_id": att.order_id,
                "status": att.status, "filled_qty": att.filled_qty, "avg_price": att.avg_price,
                "terminal": att.terminal, "error": att.error, "phase": label,
            })
            if att.terminal != "FULL_FILL":
                status_map[code] = f"FAIL:{att.terminal}"
                summary["error"] = f"{label}_not_full:{code}:{att.terminal}"
                return False
            status_map[code] = "FILLED"
        return True

    if sells:
        summary["phase"] = "SELL"
        if not send_batch(sells, "SELL"):
            return persist_failure(5)

    if buys:
        summary["phase"] = "gates_B"
        if sells:
            try:
                from ops.toss_portfolio import fetch_portfolio
                book2 = fetch_portfolio()
                account_reasons = validate_account_identity(book2, client)
                if account_reasons:
                    raise RuntimeError(",".join(account_reasons))
                cash1 = float(book2.get("cash_buying_power_usd") if book2.get("cash_buying_power_usd") is not None else book2.get("cash_usd"))
            except Exception as exc:
                summary["error"] = f"post_sell_refresh_fail:{type(exc).__name__}:{exc}"
                return persist_failure(6)
        else:
            cash1 = cash0
        gate_b = evaluate_gates(
            sendable, health, phase="B", cash_usd=cash0,
            post_sell_cash_usd=cash1, confirm_high_value=confirm_high_value,
        )
        summary["gate_b"] = gate_result_dict(gate_b)
        if not gate_b.ok:
            summary["error"] = "gate_B_fail"
            return persist_failure(7)
        summary["phase"] = "BUY"
        if not send_batch(buys, "BUY"):
            return persist_failure(5)

    summary["http_order_posts"] = getattr(client, "http_post_count", 0)
    summary["fills"] = str(fills_path)
    apply_execute_status_merge(run_dir, status_map)
    write_attempt_result(run_dir, {
        "status": "DONE", "http_order_posts": summary["http_order_posts"],
        "package_id": package_id,
    })
    _refresh_excel_after_execution(
        run_dir, positions=positions, prior_book=book, summary=summary, no_xlsx=no_xlsx,
    )
    summary["phase"] = "done"
    summary["ok"] = True
    return 0, summary
