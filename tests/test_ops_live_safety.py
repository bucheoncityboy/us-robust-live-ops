from __future__ import annotations

import json
from datetime import date, datetime
from pathlib import Path
from unittest.mock import MagicMock, patch

import pandas as pd
import pytest


# Hermetic close-panel fixture (no live data/us parquet needed on a fresh clone).
TESTS_DIR = Path(__file__).resolve().parent
PRICES_FIXTURE = TESTS_DIR / "fixtures" / "us_prices_fixture.parquet"


@pytest.fixture(autouse=True)
def no_real_http(monkeypatch):
    import requests

    def blocked(*args, **kwargs):
        raise AssertionError("real HTTP is forbidden in live-safety tests")

    for name in ("get", "post", "put", "delete", "patch"):
        monkeypatch.setattr(requests, name, blocked)
    monkeypatch.setattr(requests.Session, "get", blocked)
    monkeypatch.setattr(requests.Session, "post", blocked)


def make_package(root: Path, signal: str = "2026-06-30") -> Path:
    from ops.ops_live_safety import write_package_manifest

    run = root / signal[:7]
    run.mkdir(parents=True)
    pd.DataFrame([{"asof": signal, "code": "AAPL", "sleeve": "leader", "rank": 1}]).to_csv(run / "signals.csv", index=False)
    pd.DataFrame([{"asof": signal, "code": "AAPL", "name": "Apple", "sleeve": "leader", "final_w": 0.1}]).to_csv(run / "target.csv", index=False)
    pd.DataFrame([{
        "asof": signal, "code": "AAPL", "name": "Apple", "side": "BUY", "action": "BUY_NEW",
        "current_w": 0.0, "target_w": 0.1, "delta_w": 0.1, "sleeve": "leader", "prio": 1,
        "qty": 1.0, "notional": 100.0, "est_px": 100.0,
    }]).to_csv(run / "orders_preview.csv", index=False)
    (run / "health.json").write_text(json.dumps({
        "signal_date": signal, "status": "OK", "reasons": [], "capital_usd": 1000.0,
        "last_dates": {"prices": signal, "index": signal, "valuation": signal},
        "lag_days": {"prices": 0, "index": 0, "valuation": 0},
        "sizing": {"cash_usd": 1000.0}, "force_intramonth": False,
    }), encoding="utf-8")
    write_package_manifest(run)
    return run


def test_current_ny_month_cannot_be_bypassed_by_asof(monkeypatch):
    from ops.ops_monthly_run import is_final_month_end_signal

    close = pd.DataFrame(index=pd.to_datetime(["2026-07-24", "2026-07-27"]))
    assert not is_final_month_end_signal(
        pd.Timestamp("2026-07-27"), close,
        asof="2026-07-31", today=pd.Timestamp("2026-07-31"),
    )


def test_prior_month_completion(monkeypatch):
    from ops.ops_monthly_run import is_final_month_end_signal

    close = pd.DataFrame(index=pd.to_datetime(["2026-06-29", "2026-06-30"]))
    assert is_final_month_end_signal(pd.Timestamp("2026-06-30"), close, today=pd.Timestamp("2026-07-01"))


def test_health_reference_universe_and_sleeve_reporting():
    from ops.ops_monthly_run import build_health

    close = pd.DataFrame({"AAPL": [1.0]}, index=pd.to_datetime(["2026-06-30"]))
    idx = pd.Series([1.0], index=close.index)
    health = build_health(
        pd.Timestamp("2026-06-30"), close, pd.DataFrame(), pd.DataFrame(), idx,
        {"leader": ["AAPL"], "mom63": ["AAPL"], "lowvol": ["AAPL"]},
        {"AAPL": 1.0}, {}, reference_date="2026-07-04", universe_n=120, universe_target=150,
    )
    assert health["lag_days"]["prices"] == 4
    assert health["reference_date"] == "2026-07-04"
    assert "universe_shortfall=120/150" in health["reasons"]
    assert health["sleeve_exposure"] == {"leader": 0.6, "mom63": 0.2, "lowvol": 0.2, "cash": 0.0}
    assert health["name_primary_exposure"]["leader"] == 1.0


def test_manifest_detects_tamper(tmp_path):
    from ops.ops_live_safety import verify_package_manifest

    run = make_package(tmp_path)
    ok, reasons, _ = verify_package_manifest(run)
    assert ok and not reasons
    (run / "target.csv").write_text("asof,code,final_w\n2026-06-30,MSFT,1\n", encoding="utf-8")
    ok, reasons, _ = verify_package_manifest(run)
    assert not ok and any("hash_mismatch" in r for r in reasons)


def test_execute_package_rejects_current_unsafe_artifact():
    from ops.ops_live_safety import validate_execute_package

    result = validate_execute_package(Path("results/ops_runs/2026-07"), now=date(2026, 7, 31))
    assert not result["ok"]
    assert any("manifest" in r or "month" in r for r in result["reasons"])


def test_package_internal_date_mismatch(tmp_path):
    from ops.ops_live_safety import validate_execute_package, write_package_manifest

    run = make_package(tmp_path)
    orders = pd.read_csv(run / "orders_preview.csv")
    orders["asof"] = "2026-06-29"
    orders.to_csv(run / "orders_preview.csv", index=False)
    result = validate_execute_package(run, now=date(2026, 7, 2), max_age_days=7)
    assert not result["ok"]
    assert any("asof_mismatch" in r for r in result["reasons"])


def test_toss_weights_include_cash():
    from ops.toss_portfolio import us_weights

    book = {"total_usd": 100.0, "us_holdings": [{"symbol_raw": "AAPL", "mkt_value": 50.0}]}
    assert us_weights(book) == {"AAPL": 0.5}


def test_share_sizing_side_uses_actual_quantity():
    from ops.ops_monthly_run import attach_share_sizes

    orders = [{"code": "AAPL", "side": "SELL", "action": "SELL_TRIM", "current_w": 1.0, "target_w": 0.8}]
    close = pd.DataFrame({"AAPL": [100.0]}, index=pd.to_datetime(["2026-06-30"]))
    out = attach_share_sizes(
        orders, close, 100.0, signal_date=pd.Timestamp("2026-06-30"),
        live_px={"AAPL": 100.0}, live_qty={"AAPL": 0.5},
    )[0]
    assert out["side"] == "BUY"
    assert out["qty"] == pytest.approx(0.3)


def test_us_buy_always_uses_bounded_order_amount():
    from ops.toss_orders import build_order_payload

    payload = build_order_payload(
        symbol="AAPL", side="BUY", quantity=1.0, client_order_id="x1", order_amount=100.0,
    )
    assert payload["orderAmount"] == 100.0
    assert "quantity" not in payload
    with pytest.raises(Exception):
        build_order_payload(symbol="AAPL", side="BUY", quantity=0.1, client_order_id="x2", order_amount=0.5)


def test_partial_fill_never_full_on_partial_or_timeout():
    from ops.toss_orders import OrderClient

    client = OrderClient()
    assert client.classify_fill(10, 5, "PARTIAL_FILLED", amount_based=True) == "PARTIAL"
    assert client.classify_fill(10, 1, "PENDING", amount_based=True) == "PARTIAL"
    assert client.classify_fill(10, 1, "FILLED", amount_based=True) == "FULL_FILL"


def test_cross_check_detects_duplicate_rows():
    from ops.ops_execute_gates import cross_check

    rows = [{"code": "AAPL", "side": "BUY", "qty": 1, "notional": 100}]
    csv = pd.DataFrame(rows + rows)
    sheet = pd.DataFrame(rows + rows)
    reasons = cross_check(sheet, csv)
    assert "sheet_duplicate_order_keys" in reasons
    assert "csv_duplicate_order_keys" in reasons


def test_order_gates_minimum_cap_and_cash_buffer():
    from ops.ops_live_safety import validate_order_gates

    rows = [{"code": "AAPL", "side": "BUY", "send_qty": 0.1, "notional": 0.5}]
    reasons, _ = validate_order_gates(rows, capital_usd=1000, cash_usd=1000, phase="A")
    assert any("buy_below_minimum" in r for r in reasons)
    rows = [{"code": "AAPL", "side": "BUY", "send_qty": 1, "notional": 990}]
    reasons, _ = validate_order_gates(rows, capital_usd=10000, cash_usd=1000, phase="B")
    assert any("cash_buffer_breach" in r for r in reasons)


def test_account_selection_requires_pin_when_multiple(monkeypatch):
    import ops.toss_orders as toss_orders

    accounts = [
        {"accountSeq": "1", "accountType": "BROKERAGE"},
        {"accountSeq": "2", "accountType": "BROKERAGE"},
    ]
    monkeypatch.setattr(toss_orders, "get_secret", lambda key, default="": "")
    with pytest.raises(Exception, match="multiple brokerage"):
        toss_orders._pick_brokerage(accounts)
    monkeypatch.setattr(toss_orders, "get_secret", lambda key, default="": "2" if key == "TOSS_ACCOUNT_SEQ" else default)
    assert toss_orders._pick_brokerage(accounts)["accountSeq"] == "2"


def test_official_next_open_window():
    from ops.ops_live_safety import validate_official_next_open

    payload = {"result": {"days": [
        {"date": "2026-07-31", "regularMarket": {"start": "2026-07-31T22:30:00+09:00", "end": "2026-08-01T05:00:00+09:00"}},
        {"date": "2026-08-03", "regularMarket": {"start": "2026-08-03T22:30:00+09:00", "end": "2026-08-04T05:00:00+09:00"}},
    ]}}
    assert validate_official_next_open(payload, "2026-07-31", now="2026-08-03T22:40:00+09:00") == []
    assert validate_official_next_open(payload, "2026-07-31", now="2026-08-03T23:30:00+09:00")


def test_risk_overlay_and_sticky_reset(tmp_path, monkeypatch):
    import ops.ops_live_safety as safety

    monkeypatch.setattr(safety, "RISK_STATE_DIR", tmp_path)
    assert sum(safety.apply_risk_weights({"AAPL": 0.6, "MSFT": 0.4}, "L1").values()) == pytest.approx(0.5)
    assert safety.apply_risk_weights({"AAPL": 1.0}, "L2") == {}
    assert safety.resolve_risk_mode("L2", "1234") == "L2"
    with pytest.raises(ValueError, match="sticky L2"):
        safety.resolve_risk_mode("NORMAL", "1234")
    safety.reset_sticky_l2("1234", "manual review complete")
    assert safety.resolve_risk_mode("NORMAL", "1234") == "NORMAL"


def test_refresh_empty_download_preserves_cache(tmp_path, monkeypatch):
    import ops.us_hybrid_backtest as us

    close_path, open_path, vol_path = tmp_path / "close.parquet", tmp_path / "open.parquet", tmp_path / "us_volume_panel.parquet"
    old = pd.DataFrame({"AAPL": range(500)}, index=pd.date_range("2024-01-01", periods=500))
    old.to_parquet(close_path); old.to_parquet(open_path); old.to_parquet(vol_path)
    before = close_path.read_bytes()
    monkeypatch.setattr(us, "CACHE_PRICES", close_path)
    monkeypatch.setattr(us, "CACHE_OPEN", open_path)
    monkeypatch.setattr(us, "DATA", tmp_path)
    monkeypatch.setattr(us, "_download_ohlcv", lambda *a, **k: (pd.DataFrame(), pd.DataFrame(), pd.DataFrame()))
    with pytest.raises(RuntimeError, match="existing cache preserved"):
        us.load_prices(["AAPL"] * 100, force=True)
    assert close_path.read_bytes() == before


def test_execute_connect_failure_is_structured_zero_order_posts(tmp_path, monkeypatch):
    import ops.ops_execute as ops_execute
    import ops.ops_monthly_run as ops_monthly_run
    import ops.toss_orders as toss_orders

    run = make_package(tmp_path)
    book = {
        "account_seq": "1", "account_no_tail": "1234", "account_type": "BROKERAGE",
        "cash_buying_power_usd": 1000.0, "cash_usd": 1000.0, "us_holdings": [],
    }
    ops_monthly_run.load_positions.last_toss_book = book
    monkeypatch.setenv("TOSS_ACCOUNT_SEQ", "1")
    monkeypatch.setattr(ops_execute, "validate_execute_package", lambda *a, **k: {"ok": True, "package_id": "pkg", "signal_date": "2026-06-30"})
    monkeypatch.setattr(ops_monthly_run, "rebuild_rebalance_ticket", lambda *a, **k: {
        "n_buy": 1, "n_sell": 0, "capital_usd": 1000.0, "cash_usd": 1000.0, "risk_mode": "NORMAL",
    })
    monkeypatch.setattr(ops_execute, "LIVE_ORDERS_ENABLED", True)
    monkeypatch.setattr(ops_execute, "research_block_reason", lambda: None)
    monkeypatch.setattr(toss_orders.OrderClient, "connect", classmethod(lambda cls: (_ for _ in ()).throw(RuntimeError("connect boom"))))
    place = MagicMock()
    monkeypatch.setattr(ops_execute, "place_and_await", place)
    code, summary = ops_execute.run_execute(run, positions="toss", no_xlsx=True, risk_mode="NORMAL")
    assert code != 0
    assert summary["phase"] == "connect"
    assert summary["http_order_posts"] == 0
    assert "connect_fail" in summary["error"]
    place.assert_not_called()
    result = json.loads((run / "execute_attempt_result.json").read_text(encoding="utf-8"))
    assert result["status"] == "CONNECT_FAIL_ZERO_POST"


def test_excel_has_no_auto_approval_or_fake_live_results(tmp_path):
    from openpyxl import load_workbook
    from ops.ops_excel_write import write_ops_xlsx

    run = make_package(tmp_path)
    book = {"status": "ok", "account_no_tail": "1234", "n_us": 0, "n_kr": 0,
            "ops_ready": True, "flat": True, "equity_usd": 0, "cash_usd": 1000,
            "cash_buying_power_usd": 1000, "total_usd": 1000, "us_holdings": []}
    out = write_ops_xlsx(run, positions="none", toss_book=book, prices_path=PRICES_FIXTURE)
    wb = load_workbook(out, data_only=True, read_only=True)
    trade_values = [str(c or "") for row in wb["05_Trade"].iter_rows(min_row=1, max_row=8, values_only=True) for c in row]
    assert "CLI_REQUIRED" in trade_values
    assert "READY_TO_SEND" in trade_values
    assert "NOT_SENT" in [str(c or "") for row in wb["05_Trade"].iter_rows(values_only=True) for c in row]
    assert "Y" not in trade_values
    assert "No synthetic live performance" in str(wb["03_Results"]["A2"].value)
    wb.close()


def test_no_real_http_fixture():
    import requests
    with pytest.raises(AssertionError, match="forbidden"):
        requests.post("https://openapi.tossinvest.com/api/v1/orders")



def test_replay_status_fill_and_claim_retry(tmp_path):
    from ops.ops_live_safety import acquire_execution_claim, execution_evidence, write_attempt_result

    run = tmp_path
    sendable = [{"code": "AAPL", "side": "BUY"}]
    (run / "execute_status.json").write_text(json.dumps({"codes": {"AAPL": "FILLED"}}), encoding="utf-8")
    assert execution_evidence(run, sendable)["blocked"]
    (run / "execute_status.json").unlink()
    (run / "fills_x.jsonl").write_text(json.dumps({"event": "FILL", "code": "AAPL", "terminal": "PARTIAL"}) + "\n", encoding="utf-8")
    assert execution_evidence(run, sendable)["blocked"]
    (run / "fills_x.jsonl").unlink()
    acquire_execution_claim(run, {"package_id": "p1"})
    with pytest.raises(FileExistsError):
        acquire_execution_claim(run, {"package_id": "p1"})
    write_attempt_result(run, {"status": "CONNECT_FAIL_ZERO_POST", "http_order_posts": 0})
    acquire_execution_claim(run, {"package_id": "p1-retry"})


def test_zero_post_gate_failure_is_auto_retryable(tmp_path):
    from ops.ops_live_safety import acquire_execution_claim, execution_evidence, write_attempt_result

    run = tmp_path
    sendable = [{"code": "AAPL", "side": "BUY"}]
    acquire_execution_claim(run, {"package_id": "p1", "phase": "CONNECTING"})
    write_attempt_result(run, {"status": "GATE_FAIL_ZERO_POST", "http_order_posts": 0})
    # zero-POST gate failures are safe: no ambiguous evidence, claim can be re-acquired
    assert not execution_evidence(run, sendable)["blocked"]
    acquire_execution_claim(run, {"package_id": "p1-retry"})


def test_review_required_status_blocks_even_with_zero_posts(tmp_path):
    from ops.ops_live_safety import acquire_execution_claim, execution_evidence, write_attempt_result

    run = tmp_path
    sendable = [{"code": "AAPL", "side": "BUY"}]
    acquire_execution_claim(run, {"package_id": "p1", "phase": "CONNECTING"})
    write_attempt_result(run, {"status": "ORDER_ATTEMPT_REVIEW_REQUIRED", "http_order_posts": 0})
    assert execution_evidence(run, sendable)["blocked"]
    with pytest.raises(FileExistsError):
        acquire_execution_claim(run, {"package_id": "p1-retry"})


def test_zero_post_status_with_posts_requires_review(tmp_path):
    from ops.ops_live_safety import acquire_execution_claim, execution_evidence, write_attempt_result

    run = tmp_path
    sendable = [{"code": "AAPL", "side": "BUY"}]
    acquire_execution_claim(run, {"package_id": "p1", "phase": "CONNECTING"})
    write_attempt_result(run, {"status": "GATE_FAIL_ZERO_POST", "http_order_posts": 1})
    assert execution_evidence(run, sendable)["blocked"]


def test_resolve_capital_usd_applies_cash_buffer():
    from ops.ops_monthly_run import resolve_capital_usd

    cap, src = resolve_capital_usd(book={"total_usd": 100.0, "cash_usd": 100.0})
    assert src == "toss:total_usd"
    assert abs(cap - 98.0) < 1e-9
    cap2, src2 = resolve_capital_usd(book={"total_usd": 0.0, "cash_usd": 50.0})
    assert src2 == "toss:total_usd"  # falls back to equity+cash under total_usd source
    assert abs(cap2 - 49.0) < 1e-9
    # explicit operator capital is not buffered (operator intent)
    cap3, src3 = resolve_capital_usd(book={"total_usd": 100.0, "cash_usd": 100.0}, capital_usd=100.0)
    assert src3 == "arg" and cap3 == 100.0


def test_weekly_recompute_returns():
    from ops.ops_daily import recompute_returns

    df = recompute_returns(pd.DataFrame({
        "date": ["2026-08-14", "2026-08-07", "2026-08-21"],
        "total_usd": [101.0, 100.0, 103.0],
        "spy_close": [101.0, 100.0, 101.5],
    }))
    assert list(df["date"]) == ["2026-08-07", "2026-08-14", "2026-08-21"]
    assert abs(df["cum_ret"].iloc[-1] - 0.03) < 1e-9
    assert abs(df["ret_1d"].iloc[1] - 0.01) < 1e-9  # week-over-week pct change
    assert abs(df["spy_cum_ret"].iloc[-1] - 0.015) < 1e-9
    assert df["ret_1d"].iloc[0] != df["ret_1d"].iloc[0]  # NaN on first row


def test_weekly_backfill_skipped_weeks():
    from ops.ops_daily import backfill_missing_weeks

    spy = pd.Series(
        [100.0, 100.5, 101.0, 101.5, 102.0, 102.5, 103.0, 103.5],
        index=pd.to_datetime(["2026-07-31", "2026-08-03", "2026-08-04", "2026-08-05",
                              "2026-08-06", "2026-08-07", "2026-08-10", "2026-08-14"]),
    )
    panel = pd.DataFrame(
        {"AAPL": [10.0, 10.1, 10.2, 10.3, 10.4, 10.5, 10.6, 10.7],
         "MSFT": [20.0, 20.1, 20.2, 20.3, 20.4, 20.5, 20.6, 20.7]},
        index=spy.index,
    )
    df = pd.DataFrame([{
        "date": "2026-07-31", "total_usd": 30.0, "equity_usd": 30.0,
        "cash_usd": 1.0, "spy_close": 100.0,
    }])
    holdings = pd.DataFrame([{"date": "2026-07-31", "code": "AAPL", "qty": 1.0},
                             {"date": "2026-07-31", "code": "MSFT", "qty": 1.0}])
    out, out_h, n = backfill_missing_weeks(df, holdings, spy, panel, "2026-08-14")
    assert n == 1  # only the missed 08-07 anchor is backfilled; intra-week days are not rows
    assert list(out["date"]) == ["2026-07-31", "2026-08-07"]
    # 08-07 anchor: AAPL 10.5 + MSFT 20.5 + cash 1.0 = 32.0
    assert abs(float(out.loc[out["date"] == "2026-08-07", "total_usd"].iloc[0]) - 32.0) < 1e-6
    assert len(out_h) == 4  # 2 original + 2 backfilled codes on the anchor


def test_weekly_anchor_friday_or_holiday_adj():
    from ops.ops_daily import _is_weekly_anchor, _weekly_anchor_map

    # normal week: Friday is the anchor
    idx = pd.to_datetime(["2026-08-03", "2026-08-04", "2026-08-05", "2026-08-06", "2026-08-07"])
    m = _weekly_anchor_map(idx)
    assert list(m.values()) == [pd.Timestamp("2026-08-07")]
    assert _is_weekly_anchor(pd.Timestamp("2026-08-07"), m)
    assert not _is_weekly_anchor(pd.Timestamp("2026-08-06"), m)
    # holiday week without a Friday trading day: anchor rolls back to Thursday
    idx2 = pd.to_datetime(["2026-08-10", "2026-08-11", "2026-08-12", "2026-08-13"])
    m2 = _weekly_anchor_map(idx2)
    assert list(m2.values()) == [pd.Timestamp("2026-08-13")]


def test_weekly_anchor_mid_week_refusal():
    from ops.ops_daily import _classify_run_day

    spy = pd.Series([100.0] * 5, index=pd.to_datetime(
        ["2026-08-07", "2026-08-10", "2026-08-11", "2026-08-12", "2026-08-14"]))
    # Wednesday before the anchor -> clean no-op naming the anchor
    reason, row_date = _classify_run_day(pd.Timestamp("2026-08-12"), spy)
    assert reason == "not_weekly_anchor_day:2026-08-12:2026-08-14" and row_date is None
    # Friday anchor -> record today
    reason, row_date = _classify_run_day(pd.Timestamp("2026-08-14"), spy)
    assert reason is None and row_date == "2026-08-14"
    # Saturday and Sunday -> record the prior Friday anchor in place
    reason, row_date = _classify_run_day(pd.Timestamp("2026-08-15"), spy)
    assert reason is None and row_date == "2026-08-14"
    reason, row_date = _classify_run_day(pd.Timestamp("2026-08-16"), spy)
    assert reason is None and row_date == "2026-08-14"
    # non-trading weekday (holiday) -> catch-all refusal
    reason, row_date = _classify_run_day(pd.Timestamp("2026-09-07"), spy)
    assert reason is not None and reason.startswith("not_a_us_trading_day_or_missing_spy")
    # missing SPY series -> catch-all refusal
    reason, row_date = _classify_run_day(pd.Timestamp("2026-08-14"), None)
    assert reason is not None and reason.startswith("not_a_us_trading_day_or_missing_spy")


def test_weekend_run_records_prior_anchor(monkeypatch):
    import ops.ops_daily as od
    import ops.ops_live_safety as ols

    spy = pd.Series([100.0, 101.0, 102.0], index=pd.to_datetime(["2026-08-07", "2026-08-10", "2026-08-14"]))
    writes = {}
    monkeypatch.setattr(od, "_spy_series", lambda: spy)
    monkeypatch.setattr(od, "load_daily", lambda: pd.DataFrame(columns=od.COLS))
    monkeypatch.setattr(od, "load_holdings", lambda: pd.DataFrame(columns=od.HOLD_COLS))
    monkeypatch.setattr(od, "load_close_panel", lambda: None)
    monkeypatch.setattr(od, "_read_csv", lambda path, cols: pd.DataFrame(columns=cols))
    monkeypatch.setattr(ols, "atomic_write_csv", lambda df, path: writes.__setitem__(str(path), df.copy()))
    book = {"total_usd": 100.0, "equity_usd": 99.0, "cash_usd": 1.0, "us_holdings": [], "qty_map": {}}

    # Saturday 2026-08-08 -> records the prior Friday anchor 08-07 in place
    monkeypatch.setattr(ols, "ny_today", lambda: pd.Timestamp("2026-08-08").date())
    res = od.record(book=book)
    assert res["recorded"] and res["date"] == "2026-08-07"
    assert abs(res["spy_close"] - 100.0) < 1e-9
    nav_writes = [v for k, v in writes.items() if k.endswith("daily_nav.csv")]
    assert nav_writes and str(nav_writes[0]["date"].iloc[0]) == "2026-08-07"

    # Sunday re-run overwrites the same anchor row (idempotent per anchor)
    monkeypatch.setattr(ols, "ny_today", lambda: pd.Timestamp("2026-08-09").date())
    res2 = od.record(book=book)
    assert res2["recorded"] and res2["date"] == "2026-08-07"

    # Monday after the missed anchor -> clean no-op; gap backfilled on the next anchor run
    monkeypatch.setattr(ols, "ny_today", lambda: pd.Timestamp("2026-08-10").date())
    res3 = od.record(book=book)
    assert not res3["recorded"] and res3["reason"].startswith("not_weekly_anchor_day:2026-08-10:2026-08-14")


def test_compress_daily_to_weekly():
    from ops.ops_daily import compress_daily_to_weekly

    spy = pd.Series(
        [99.0, 100.0, 101.0, 102.0, 103.0, 104.0, 105.0],
        index=pd.to_datetime(["2026-07-31", "2026-08-03", "2026-08-04", "2026-08-05",
                              "2026-08-06", "2026-08-07", "2026-08-10"]),
    )
    panel = pd.DataFrame(
        {"AAPL": [10.0, 10.2, 10.4, 10.6, 11.0],
         "MSFT": [20.0, 20.1, 20.2, 20.3, 19.5]},
        index=pd.to_datetime(["2026-08-03", "2026-08-04", "2026-08-05", "2026-08-06", "2026-08-07"]),
    )
    df = pd.DataFrame([
        {"date": "2026-08-03", "total_usd": 31.0, "equity_usd": 30.0, "cash_usd": 1.0, "spy_close": 100.0},
        {"date": "2026-08-04", "total_usd": 31.3, "equity_usd": 30.3, "cash_usd": 1.0, "spy_close": 101.0},
        {"date": "2026-08-05", "total_usd": 31.6, "equity_usd": 30.6, "cash_usd": 1.0, "spy_close": 102.0},
        {"date": "2026-08-06", "total_usd": 31.9, "equity_usd": 30.9, "cash_usd": 1.0, "spy_close": 103.0},
    ])
    holdings = pd.DataFrame([
        {"date": "2026-08-06", "code": "AAPL", "qty": 1.0, "weight": 0.346, "avg_cost": 10.0, "sleeve": "leader"},
        {"date": "2026-08-06", "code": "MSFT", "qty": 1.0, "weight": 0.654, "avg_cost": 20.0, "sleeve": "mom63"},
    ])
    out, out_h, n = compress_daily_to_weekly(df, holdings, spy, panel)
    assert n == 1
    assert list(out["date"]) == ["2026-08-07"]  # 08-07 Friday close in the fixture panel
    # anchor: AAPL 11.0 + MSFT 19.5 + cash 1.0 = 31.5; seeded ret_1d = 31.5 / 31.0 - 1
    assert abs(float(out["total_usd"].iloc[0]) - 31.5) < 1e-6
    assert abs(float(out["ret_1d"].iloc[0]) - (31.5 / 31.0 - 1.0)) < 1e-6  # stored rounded to 6 dp
    assert abs(float(out["cum_ret"].iloc[0])) < 1e-12
    assert set(out_h["code"]) == {"AAPL", "MSFT"}
    assert (out_h["date"] == "2026-08-07").all()
    # idempotent: fully compressed input passes through unchanged
    out2, _, n2 = compress_daily_to_weekly(out, out_h, spy, panel)
    assert n2 == 0 and list(out2["date"]) == ["2026-08-07"]
def test_record_write_cleans_stale_legacy_dates(tmp_path, monkeypatch):
    import ops.ops_daily as od
    import ops.ops_live_safety as ols

    spy = pd.Series(
        [100.0, 100.5, 101.0, 101.5, 102.0, 103.0],
        index=pd.to_datetime(["2026-08-03", "2026-08-04", "2026-08-05", "2026-08-06", "2026-08-07", "2026-08-10"]),
    )
    writes = {}
    # realistic pre-migration state: the union and the on-disk month file hold legacy
    # daily rows (08-03..08-06) and the close panel covers the 08-07 anchor
    nav_union = pd.DataFrame([
        {"date": "2026-08-03", "total_usd": 31.0, "equity_usd": 30.0, "cash_usd": 1.0, "spy_close": 100.0},
        {"date": "2026-08-04", "total_usd": 31.3, "equity_usd": 30.3, "cash_usd": 1.0, "spy_close": 100.5},
        {"date": "2026-08-05", "total_usd": 31.6, "equity_usd": 30.6, "cash_usd": 1.0, "spy_close": 101.0},
        {"date": "2026-08-06", "total_usd": 31.9, "equity_usd": 30.9, "cash_usd": 1.0, "spy_close": 101.5},
    ])
    hold_union = pd.DataFrame([
        {"date": "2026-08-06", "code": "AAPL", "qty": 1.0, "weight": 0.5, "avg_cost": 10.0, "sleeve": "leader"},
        {"date": "2026-08-06", "code": "MSFT", "qty": 1.0, "weight": 0.5, "avg_cost": 20.0, "sleeve": "mom63"},
    ])
    panel = pd.DataFrame(
        {"AAPL": [10.0, 10.2, 10.4, 10.6, 11.0],
         "MSFT": [20.0, 20.1, 20.2, 20.3, 19.5]},
        index=pd.to_datetime(["2026-08-03", "2026-08-04", "2026-08-05", "2026-08-06", "2026-08-07"]),
    )
    monkeypatch.setattr(od, "_spy_series", lambda: spy)
    monkeypatch.setattr(od, "load_daily", lambda: nav_union.copy())
    monkeypatch.setattr(od, "load_holdings", lambda: hold_union.copy())
    monkeypatch.setattr(od, "load_close_panel", lambda: panel)
    monkeypatch.setattr(od, "_daily_path", lambda month: Path(tmp_path) / f"daily_nav_{month}.csv")
    monkeypatch.setattr(od, "_holdings_path", lambda month: Path(tmp_path) / f"daily_holdings_{month}.csv")
    monkeypatch.setattr(od, "_read_csv", lambda path, cols: nav_union.copy() if str(path).endswith("daily_nav.csv") else hold_union.copy())
    monkeypatch.setattr(ols, "atomic_write_csv", lambda df, path: writes.__setitem__(str(path), df.copy()))
    monkeypatch.setattr(ols, "ny_today", lambda: pd.Timestamp("2026-08-08").date())  # Saturday
    book = {"total_usd": 100.0, "equity_usd": 99.0, "cash_usd": 1.0,
            "us_holdings": [{"symbol_raw": "AAPL", "mkt_value": 50.0, "avg_price": 10.0},
                            {"symbol_raw": "MSFT", "mkt_value": 49.0, "avg_price": 20.0}],
            "qty_map": {"AAPL": 1.0, "MSFT": 1.0}}
    res = od.record(book=book)
    assert res["recorded"] and res["date"] == "2026-08-07"
    nav_write = [v for k, v in writes.items() if "daily_nav" in k][0]
    # the legacy daily rows are compressed to the anchor and the written month file
    # contains only the anchor row (stale dates cleaned)
    assert list(nav_write["date"]) == ["2026-08-07"]
    hold_write = [v for k, v in writes.items() if "daily_holdings" in k][0]
    assert set(hold_write["date"]) == {"2026-08-07"}
    # integrated seed preservation: the first anchor's ret_1d is the real first-week
    # return (live anchor total / legacy first recorded total - 1), re-applied after
    # the final recompute instead of being wiped
    first_anchor_total = float(nav_write["total_usd"].iloc[0])
    assert abs(float(nav_write["ret_1d"].iloc[0]) - (first_anchor_total / 31.0 - 1.0)) < 1e-6


def test_excel_stock_weekly_stepping():
    from ops.ops_daily import prior_anchor as _prior_anchor

    dates = ["2026-08-07", "2026-08-14", "2026-08-21"]
    assert _prior_anchor("2026-08-07", dates) is None
    assert _prior_anchor("2026-08-14", dates) == "2026-08-07"
    assert _prior_anchor("2026-08-21", dates) == "2026-08-14"
    # anchor-to-anchor semantics: the panel holds daily closes incl. 08-13 (Thu);
    # the weekly return for 08-14 must use the 08-07 anchor close, NOT the prior trading day
    panel = pd.DataFrame(
        {"AAPL": [100.0, 100.5, 101.5, 101.6, 102.0]},
        index=pd.to_datetime(["2026-08-07", "2026-08-13", "2026-08-14", "2026-08-20", "2026-08-21"]),
    )
    prior = _prior_anchor("2026-08-14", dates)
    ret = panel.at[pd.Timestamp("2026-08-14"), "AAPL"] / panel.at[pd.Timestamp(prior), "AAPL"] - 1.0
    assert abs(ret - (101.5 / 100.0 - 1.0)) < 1e-9  # 08-07 -> 08-14, NOT 08-13 -> 08-14


def test_weekly_assets_script_smoke():
    """The weekly assets script must not crash on the legacy-ledger guard path (import/call-site binding)."""
    import subprocess
    import sys

    from collections import Counter

    from ops.ops_daily import load_daily

    nav = load_daily()
    week_counts = Counter()
    for d in nav["date"]:
        iso = pd.Timestamp(d).isocalendar()
        week_counts[(iso.year, iso.week)] += 1
    if not any(n > 1 for n in week_counts.values()):
        pytest.skip("ledger already migrated; script guard path not exercisable")
    script = Path(__file__).resolve().parent.parent / "scripts" / "make_weekly_notion_assets.py"
    r = subprocess.run(
        [sys.executable, str(script)],
        capture_output=True, text=True, timeout=120,
    )
    assert r.returncode == 1, f"expected guard exit 1, got {r.returncode}\n{r.stdout}\n{r.stderr}"
    assert "WARNING" in r.stdout
    assert "NameError" not in r.stderr and "Traceback" not in r.stderr

def test_compress_mixed_deferred_and_synthesized_weeks():
    """Partial migration: a deferred (unpriccable) legacy week followed by a priced week
    must not crash on mixed Series/dict rows and must keep both weeks' rows."""
    from ops.ops_daily import compress_daily_to_weekly

    spy = pd.Series(
        [100.0, 101.0, 102.0, 103.0, 104.0, 105.0, 106.0, 107.0, 108.0],
        index=pd.to_datetime(["2026-07-31", "2026-08-03", "2026-08-04", "2026-08-05",
                              "2026-08-06", "2026-08-07", "2026-08-10", "2026-08-11", "2026-08-14"]),
    )
    # panel covers only the SECOND week's anchor (08-14); the first week's anchor
    # (08-07) is missing from the panel -> that week is deferred (rows kept as-is)
    panel = pd.DataFrame(
        {"AAPL": [10.0, 10.2, 10.4, 10.6, 11.0]},
        index=pd.to_datetime(["2026-08-10", "2026-08-11", "2026-08-12", "2026-08-13", "2026-08-14"]),
    )
    df = pd.DataFrame([
        {"date": "2026-08-03", "total_usd": 31.0, "equity_usd": 30.0, "cash_usd": 1.0, "spy_close": 100.0},
        {"date": "2026-08-04", "total_usd": 31.3, "equity_usd": 30.3, "cash_usd": 1.0, "spy_close": 101.0},
        {"date": "2026-08-05", "total_usd": 31.6, "equity_usd": 30.6, "cash_usd": 1.0, "spy_close": 102.0},
        {"date": "2026-08-06", "total_usd": 31.9, "equity_usd": 30.9, "cash_usd": 1.0, "spy_close": 103.0},
        {"date": "2026-08-10", "total_usd": 32.0, "equity_usd": 31.0, "cash_usd": 1.0, "spy_close": 105.0},
        {"date": "2026-08-11", "total_usd": 32.4, "equity_usd": 31.4, "cash_usd": 1.0, "spy_close": 106.0},
    ])
    holdings = pd.DataFrame([
        {"date": "2026-08-06", "code": "AAPL", "qty": 1.0, "weight": 1.0, "avg_cost": 10.0, "sleeve": "leader"},
        {"date": "2026-08-11", "code": "AAPL", "qty": 1.0, "weight": 1.0, "avg_cost": 10.0, "sleeve": "leader"},
    ])
    out, out_h, n = compress_daily_to_weekly(df, holdings, spy, panel)
    assert n == 1  # only the second week (08-14 anchor) is synthesized
    assert list(out["date"]) == ["2026-08-03", "2026-08-04", "2026-08-05", "2026-08-06", "2026-08-14"]
    # second-week anchor: AAPL 11.0 + cash 1.0 = 12.0; seeded from 32.0 (week's first total)
    assert abs(float(out.loc[out["date"] == "2026-08-14", "total_usd"].iloc[0]) - 12.0) < 1e-6
    assert abs(float(out.loc[out["date"] == "2026-08-14", "ret_1d"].iloc[0]) - (12.0 / 32.0 - 1.0)) < 1e-6
    assert set(out_h["date"]) == {"2026-08-06", "2026-08-14"}

def test_attach_share_sizes_holds_micro_trims():
    from ops.ops_monthly_run import attach_share_sizes

    close = pd.DataFrame({"DDOG": [274.17]}, index=pd.to_datetime(["2026-08-03"]))
    # held 0.000037 shares above target (=$0.01): must HOLD, not SELL_TRIM
    orders = [{"code": "DDOG", "current_w": 0.0785, "target_w": 0.08, "side": "HOLD", "action": "HOLD"}]
    out = attach_share_sizes(
        orders, close, 69.11,
        live_px={"DDOG": 274.17}, live_qty={"DDOG": 0.020202},
        signal_date=pd.Timestamp("2026-07-31"),
    )
    assert out[0]["side"] == "HOLD" and out[0]["qty"] == 0.0
    # real trim (~$8) must still SELL
    orders2 = [{"code": "DDOG", "current_w": 0.10, "target_w": 0.08, "side": "SELL", "action": "SELL_TRIM"}]
    out2 = attach_share_sizes(
        orders2, close, 69.11,
        live_px={"DDOG": 274.17}, live_qty={"DDOG": 0.05},
        signal_date=pd.Timestamp("2026-07-31"),
    )
    assert out2[0]["side"] == "SELL" and float(out2[0]["qty"]) > 0


def test_official_calendar_start_time_schema():
    from ops.ops_live_safety import validate_official_next_open

    payload = {"result": {
        "today": {"date": "2026-07-31", "regularMarket": {"startTime": "2026-07-31T22:30:00+09:00", "endTime": "2026-08-01T05:00:00+09:00"}},
        "nextBusinessDay": {"date": "2026-08-03", "regularMarket": {"startTime": "2026-08-03T22:30:00+09:00", "endTime": "2026-08-04T05:00:00+09:00"}},
    }}
    assert validate_official_next_open(payload, "2026-07-31", now="2026-08-03T22:35:00+09:00") == []


def test_high_value_requires_explicit_confirmation():
    from ops.ops_live_safety import validate_order_gates

    row = [{"code": "AAPL", "side": "BUY", "send_qty": 1, "notional": 70000.0}]
    reasons, _ = validate_order_gates(row, capital_usd=1_000_000, cash_usd=1_000_000, phase="A")
    assert any("high_value_confirmation_required" in r for r in reasons)
    reasons, _ = validate_order_gates(row, capital_usd=1_000_000, cash_usd=1_000_000, phase="A", confirm_high_value=True)
    assert not any("high_value_confirmation_required" in r for r in reasons)



def test_parquet_month_end_readiness_uses_actual_index(tmp_path):
    from datetime import timezone
    from ops.ops_excel_write import compute_month_end_readiness

    prices = tmp_path / "prices.parquet"
    pd.DataFrame({"AAPL": [1.0, 2.0]}, index=pd.to_datetime(["2026-07-28", "2026-07-29"])).to_parquet(prices)
    result = compute_month_end_readiness(
        tmp_path / "2026-07",
        {"signal_date": "2026-07-23"},
        prices_path=prices,
        now=datetime(2026, 7, 30, 17, 0, tzinfo=timezone.utc),
    )
    assert result["expected_close"] == "2026-07-31"
    assert result["actual_prices"] == "2026-07-29"
    assert result["package_signal"] == "2026-07-23"
    assert result["status"] == "BLOCK"
    assert result["reasons"] == ["WAIT_FOR_CLOSE", "PRICE_BEHIND", "PACKAGE_SIGNAL_MISMATCH"]
    assert result["checked_kst"].endswith("KST")


def test_execution_status_dict_is_normalized():
    from ops.ops_excel_write import _normalize_execution_status

    assert _normalize_execution_status({"status": "FILLED", "terminal": "FULL_FILL"}) == "FILLED"
    assert _normalize_execution_status({"terminal": "PARTIAL"}) == "FAIL:PARTIAL"
    assert _normalize_execution_status({"status": "TIMEOUT_OPEN"}) == "FAIL:TIMEOUT_OPEN"
    assert _normalize_execution_status(None) == "NOT_SENT"


def test_execution_results_accumulate_partial_retry(tmp_path):
    from ops.ops_excel_write import _execution_results

    orders = pd.DataFrame([{"code": "AAPL", "side": "BUY", "qty": 10.0, "est_px": 100.0}])
    first = {
        "event": "FILL", "utc": "2026-07-31T13:31:00+00:00", "code": "AAPL", "side": "BUY",
        "qty": 10.0, "filled_qty": 4.0, "avg_price": 100.0, "status": "PARTIAL_FILLED",
        "terminal": "PARTIAL", "order_id": "o1", "client_order_id": "c1",
    }
    second = {
        "event": "FILL", "utc": "2026-07-31T13:32:00+00:00", "code": "AAPL", "side": "BUY",
        "qty": 6.0, "filled_qty": 6.0, "avg_price": 102.0, "status": "FILLED",
        "terminal": "FULL_FILL", "order_id": "o2", "client_order_id": "c2",
    }
    (tmp_path / "fills_1.jsonl").write_text(json.dumps(first) + "\n", encoding="utf-8")
    (tmp_path / "fills_2.jsonl").write_text(json.dumps(second) + "\n", encoding="utf-8")
    (tmp_path / "execute_status.json").write_text(json.dumps({
        "codes": {"AAPL": {"status": "FILLED", "terminal": "FULL_FILL"}}
    }), encoding="utf-8")
    rows = _execution_results(tmp_path, orders)
    assert len(rows) == 1
    row = rows[0]
    assert row["requested_qty"] == pytest.approx(10.0)
    assert row["filled_qty"] == pytest.approx(10.0)
    assert row["avg_price"] == pytest.approx(101.2)
    assert row["adverse_slippage"] == pytest.approx(0.012)
    assert row["status"] == "FILLED"
    assert row["order_ids_text"] == "o1\no2"
    assert row["client_order_ids_text"] == "c1\nc2"


def test_excel_shows_month_end_block_execution_results_and_stale_book(tmp_path, monkeypatch):
    from openpyxl import load_workbook
    import ops.ops_excel_write as excel

    run = make_package(tmp_path, signal="2026-07-23")
    monkeypatch.setattr(excel, "compute_month_end_readiness", lambda *a, **k: {
        "expected_close": "2026-07-31", "actual_prices": "2026-07-29",
        "package_signal": "2026-07-23", "status": "BLOCK",
        "checked_kst": "2026-07-31 02:11:03 KST",
        "reasons": ["WAIT_FOR_CLOSE", "PRICE_BEHIND", "PACKAGE_SIGNAL_MISMATCH"],
        "reason": "WAIT_FOR_CLOSE | PRICE_BEHIND | PACKAGE_SIGNAL_MISMATCH",
    })
    fill = {
        "event": "FILL", "utc": "2026-07-30T17:00:00+00:00", "code": "AAPL", "side": "BUY",
        "qty": 1.0, "filled_qty": 0.4, "avg_price": 101.0, "status": "PARTIAL_FILLED",
        "terminal": "PARTIAL", "order_id": "long-order-id", "client_order_id": "long-client-id",
    }
    (run / "fills_x.jsonl").write_text(json.dumps(fill) + "\n", encoding="utf-8")
    (run / "execute_status.json").write_text(json.dumps({"codes": {
        "AAPL": {"status": "PARTIAL_FILLED", "terminal": "PARTIAL"}
    }}), encoding="utf-8")
    stale_book = {
        "status": "stale", "_stale": True, "error": "refresh timeout", "account_no_tail": "1234",
        "n_us": 1, "n_kr": 0, "ops_ready": False, "flat": False, "cash_usd": 500.0,
        "cash_buying_power_usd": 500.0, "equity_usd": 500.0, "total_usd": 1000.0,
        "us_holdings": [{"symbol_raw": "AAPL", "name": "Apple", "qty": 5, "avg_price": 90,
                         "last_price": 100, "mkt_value": 500, "weight": 0.5}],
    }
    out = excel.write_ops_xlsx(run, positions="none", toss_book=stale_book)
    wb = load_workbook(out, data_only=True, read_only=True)
    health_values = [str(c or "") for row in wb["04_Health"].iter_rows(values_only=True) for c in row]
    assert all(value in health_values for value in [
        "Expected Close", "2026-07-31", "Actual Prices", "2026-07-29",
        "Package Signal", "2026-07-23", "Month-End Ready", "BLOCK",
        "WAIT_FOR_CLOSE | PRICE_BEHIND | PACKAGE_SIGNAL_MISMATCH",
    ])
    trade = wb["05_Trade"]
    trade_values = [str(c or "") for row in trade.iter_rows(values_only=True) for c in row]
    assert "MONTH_END_BLOCK" in trade_values
    assert "MONTH_END_NOT_READY" in trade_values
    assert "FAIL:PARTIAL" in trade_values
    assert "C. EXECUTION RESULTS (cumulative across fills_*.jsonl)" in trade_values
    assert "long-order-id" in trade_values and "long-client-id" in trade_values
    assert trade.cell(10, 20).value == "reason"
    portfolio_values = [str(c or "") for row in wb["01_Portfolio_Now"].iter_rows(values_only=True) for c in row]
    assert "STALE / REFRESH ERROR" in portfolio_values
    assert "AAPL" in portfolio_values
    assert "BLOCK" in portfolio_values
    wb.close()


def test_workbook_publish_recovery_then_main_cleanup(tmp_path, monkeypatch):
    import os
    import ops.ops_excel_write as excel

    main = tmp_path / "US_Robust_Ops_2026-07.xlsx"
    updated = tmp_path / "US_Robust_Ops_2026-07_updated.xlsx"
    tmp = tmp_path / ".~tmp_US_Robust_Ops_2026-07.xlsx"
    tmp.write_text("new", encoding="utf-8")
    main.write_text("locked-main", encoding="utf-8")
    updated.write_text("locked-updated", encoding="utf-8")
    real_replace = os.replace

    def locked_replace(src, dst):
        if Path(dst) in {main, updated}:
            raise PermissionError("locked")
        return real_replace(src, dst)

    monkeypatch.setattr(excel.os, "replace", locked_replace)
    recovery = excel._publish_workbook(tmp, main, now=datetime(2026, 7, 31, 2, 11, 3))
    assert recovery.name == "US_Robust_Ops_2026-07_recovery_20260731_021103.xlsx"
    assert recovery.read_text(encoding="utf-8") == "new"

    monkeypatch.setattr(excel.os, "replace", real_replace)
    tmp2 = tmp_path / ".~tmp_second.xlsx"
    tmp2.write_text("latest", encoding="utf-8")
    out = excel._publish_workbook(tmp2, main)
    assert out == main and main.read_text(encoding="utf-8") == "latest"
    assert not updated.exists()
    assert not list(tmp_path.glob("US_Robust_Ops_2026-07_recovery_*.xlsx"))
    assert not list(tmp_path.glob(".~tmp_*"))


def test_find_xlsx_selects_latest_mtime(tmp_path):
    import os
    from ops.ops_execute import _find_xlsx

    main = tmp_path / "US_Robust_Ops_2026-07.xlsx"
    recovery = tmp_path / "US_Robust_Ops_2026-07_recovery_20260731_021103.xlsx"
    main.write_text("old", encoding="utf-8")
    recovery.write_text("new", encoding="utf-8")
    os.utime(main, (1, 1))
    os.utime(recovery, (2, 2))
    assert _find_xlsx(tmp_path) == recovery


def test_post_trade_refresh_failure_preserves_prior_book_as_stale(tmp_path, monkeypatch):
    import ops.ops_execute as ops_execute
    import ops.ops_excel_write as ops_excel_write
    import ops.toss_portfolio as toss_portfolio

    captured = {}
    prior = {"us_holdings": [{"symbol_raw": "AAPL", "qty": 2}], "cash_usd": 10.0, "n_us": 1}
    monkeypatch.setattr(toss_portfolio, "fetch_portfolio", lambda: (_ for _ in ()).throw(RuntimeError("timeout")))

    def fake_write(run_dir, **kwargs):
        captured["book"] = kwargs["toss_book"]
        return Path(run_dir) / "US_Robust_Ops_2026-07_updated.xlsx"

    monkeypatch.setattr(ops_excel_write, "write_ops_xlsx", fake_write)
    summary = {}
    out = ops_execute._refresh_excel_after_execution(
        tmp_path, positions="toss", prior_book=prior, summary=summary, no_xlsx=False,
    )
    assert out.name.endswith("_updated.xlsx")
    assert captured["book"]["_stale"] is True
    assert captured["book"]["us_holdings"] == prior["us_holdings"]
    assert captured["book"]["ops_ready"] is False
    assert summary["post_trade_portfolio"] == "STALE / REFRESH ERROR"


@pytest.mark.parametrize("terminal,expected_code", [("PARTIAL", 5), ("FULL_FILL", 0)])
def test_execute_order_outcomes_rebuild_post_trade_excel(tmp_path, monkeypatch, terminal, expected_code):
    from types import SimpleNamespace
    import ops.ops_execute as ops_execute
    import ops.ops_monthly_run as ops_monthly_run
    import ops.ops_excel_write as ops_excel_write
    import ops.toss_orders as toss_orders

    run = make_package(tmp_path / terminal.lower())
    book = {
        "account_seq": "1", "account_no_tail": "1234", "account_type": "BROKERAGE",
        "cash_buying_power_usd": 1000.0, "cash_usd": 1000.0, "us_holdings": [],
    }
    ops_monthly_run.load_positions.last_toss_book = book
    monkeypatch.setenv("TOSS_ACCOUNT_SEQ", "1")
    monkeypatch.setattr(ops_execute, "validate_execute_package", lambda *a, **k: {
        "ok": True, "package_id": f"pkg-{terminal}", "signal_date": "2026-06-30"
    })
    monkeypatch.setattr(ops_monthly_run, "rebuild_rebalance_ticket", lambda *a, **k: {
        "n_buy": 1, "n_sell": 0, "capital_usd": 1000.0, "cash_usd": 1000.0, "risk_mode": "NORMAL",
    })
    monkeypatch.setattr(ops_execute, "LIVE_ORDERS_ENABLED", True)
    monkeypatch.setattr(ops_execute, "research_block_reason", lambda: None)
    monkeypatch.setattr(ops_execute, "validate_account_identity", lambda *a, **k: [])
    monkeypatch.setattr(ops_execute, "validate_official_next_open", lambda *a, **k: [])
    monkeypatch.setattr(ops_excel_write, "write_ops_xlsx", lambda *a, **k: run / "not-created.xlsx")

    class Client:
        http_post_count = 1
        account_seq = "1"
        def get_us_market_calendar(self, signal):
            return {}
        def get_open_orders(self):
            return []

    monkeypatch.setattr(toss_orders.OrderClient, "connect", classmethod(lambda cls: Client()))
    monkeypatch.setattr(ops_execute, "place_and_await", lambda *a, **k: SimpleNamespace(
        client_order_id="cid", order_id="oid", status="FILLED" if terminal == "FULL_FILL" else "PARTIAL_FILLED",
        filled_qty=1.0 if terminal == "FULL_FILL" else 0.4, avg_price=100.0,
        terminal=terminal, error=None,
    ))
    refresh = MagicMock(return_value=run / "post.xlsx")
    monkeypatch.setattr(ops_execute, "_refresh_excel_after_execution", refresh)
    code, _ = ops_execute.run_execute(run, positions="toss", no_xlsx=False, risk_mode="NORMAL")
    assert code == expected_code
    refresh.assert_called_once()
